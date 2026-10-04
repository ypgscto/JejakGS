# -*- coding: utf-8 -*-
"""Generate formal JejakGS alumni user guide as Word document."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

OUTPUT = r"C:\laragon\www\jejakgs\docs\Panduan_Penggunaan_JejakGS_Alumni.docx"

MAROON = RGBColor(0x80, 0x00, 0x20)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x44, 0x44, 0x44)


def set_run_font(run, size=11, bold=False, color=BLACK, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_paragraph(
    doc,
    text,
    *,
    size=11,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_after=8,
    space_before=0,
    first_line_indent=None,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line_indent is not None:
        p.paragraph_format.first_line_indent = first_line_indent
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    size = 14 if level == 1 else 12
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=MAROON)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        # Clear default runs and set font
        if p.runs:
            for r in p.runs:
                r.text = ""
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        if p.runs:
            for r in p.runs:
                r.text = ""
        run = p.add_run(item)
        set_run_font(run, size=11)


def set_cell_shading(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(cell, "800020")

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_run_font(run, size=10)

    doc.add_paragraph()
    return table


def add_horizontal_line(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "800020")
    pBdr.append(bottom)
    pPr.append(pBdr)


def build():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

    # --- Cover ---
    for _ in range(3):
        doc.add_paragraph()

    add_paragraph(
        doc,
        "SEKOLAH TINGGI ILMU KESEHATAN GUNUNG SARI",
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_paragraph(
        doc,
        "STIKES Gunung Sari",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )
    add_horizontal_line(doc)

    add_paragraph(
        doc,
        "PANDUAN PENGGUNAAN",
        size=18,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_paragraph(
        doc,
        "APLIKASI JEJAKGS",
        size=18,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_paragraph(
        doc,
        "Bagi Alumni STIKES Gunung Sari",
        size=13,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=24,
    )

    add_paragraph(
        doc,
        "Dokumen ini disusun sebagai pedoman resmi penggunaan aplikasi JejakGS "
        "bagi alumni dalam mengakses layanan alumni digital STIKES Gunung Sari.",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=36,
    )

    add_paragraph(
        doc,
        "Tahun 2026",
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=0,
    )

    doc.add_page_break()

    # --- Intro ---
    add_heading_custom(doc, "I. PENDAHULUAN", 1)
    add_paragraph(
        doc,
        "JejakGS merupakan aplikasi mobile resmi bagi alumni STIKES Gunung Sari. "
        "Aplikasi ini menyediakan layanan alumni secara digital, meliputi pendaftaran "
        "dan verifikasi alumni, pengisian Tracer Study, informasi lowongan kerja, "
        "pengumuman, kegiatan alumni, kartu alumni digital, jejak angkatan, serta "
        "layanan Ikatan Alumni (IKA).",
        first_line_indent=Cm(1),
    )
    add_paragraph(
        doc,
        "Penting untuk dipahami bahwa status alumni terverifikasi berbeda dengan "
        "status anggota IKA. Setelah akun alumni diverifikasi, pengguna dapat "
        "mengakses fitur layanan alumni. Adapun fitur khusus anggota IKA—seperti "
        "kartu anggota IKA, iuran atau donasi, voting, forum, dan event anggota—"
        "hanya dapat digunakan setelah alumni mendaftar IKA dan disetujui sebagai "
        "anggota aktif.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "II. PERSIAPAN", 1)
    add_paragraph(
        doc,
        "Sebelum menggunakan aplikasi JejakGS, alumni diharapkan menyiapkan hal-hal berikut:",
    )
    add_numbered(
        doc,
        [
            "Perangkat smartphone Android yang terhubung dengan internet.",
            "Aplikasi JejakGS yang diunduh melalui saluran resmi (Google Play Store atau distribusi resmi kampus).",
            "Alamat email aktif.",
            "Nomor Induk Mahasiswa (NIM).",
            "Foto ijazah yang jelas dan terbaca untuk keperluan verifikasi.",
            "Foto diri terbaru.",
        ],
    )

    add_heading_custom(doc, "III. PENDAFTARAN DAN MASUK APLIKASI", 1)
    add_heading_custom(doc, "3.1 Pendaftaran Akun", 2)
    add_numbered(
        doc,
        [
            "Buka aplikasi JejakGS.",
            "Pilih menu Register Alumni.",
            "Isi data email, nama lengkap, kata sandi, konfirmasi kata sandi, dan NIM.",
            "Kirim formulir pendaftaran.",
            "Apabila diminta, lakukan konfirmasi email melalui tautan yang dikirim ke alamat email.",
            "Kembali ke aplikasi dan lakukan masuk (login).",
        ],
    )

    add_heading_custom(doc, "3.2 Masuk (Login)", 2)
    add_numbered(
        doc,
        [
            "Buka aplikasi JejakGS.",
            "Masukkan email dan kata sandi.",
            "Ketuk tombol Masuk.",
        ],
    )
    add_paragraph(
        doc,
        "Apabila muncul pemberitahuan bahwa email belum dikonfirmasi, lakukan "
        "pengiriman ulang konfirmasi email, kemudian coba masuk kembali.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "3.3 Keluar (Logout)", 2)
    add_paragraph(
        doc,
        "Untuk keluar dari aplikasi, buka tab Profil atau layar status verifikasi, "
        "kemudian ketuk tombol Keluar.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "IV. PELENGKAPAN PROFIL DAN VERIFIKASI ALUMNI", 1)
    add_paragraph(
        doc,
        "Pada penggunaan awal, alumni diwajibkan melengkapi data profil dan "
        "mengirimkan berkas pendukung sebelum dapat mengakses menu utama.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "4.1 Melengkapi Profil", 2)
    add_numbered(
        doc,
        [
            "Isi NIM, nama, tahun angkatan, dan tahun lulus.",
            "Unggah foto ijazah.",
            "Unggah foto diri terbaru.",
            "Ketuk tombol Kirim untuk Verifikasi.",
        ],
    )
    add_paragraph(doc, "Ketentuan foto ijazah:")
    add_bullets(
        doc,
        [
            "Diambil dengan pencahayaan yang cukup.",
            "Nama, NIM, dan data penting harus terbaca dengan jelas.",
            "Tidak buram dan tidak terpotong.",
            "Foto ijazah hanya digunakan untuk pemeriksaan oleh admin dan tidak ditampilkan kepada alumni lain.",
        ],
    )

    add_heading_custom(doc, "4.2 Proses Verifikasi", 2)
    add_paragraph(
        doc,
        "Setelah data dikirim, pengajuan akan diperiksa oleh admin. Pada layar status, "
        "alumni dapat melihat ringkasan data, memperbarui status melalui tombol "
        "Cek Status Terbaru, memperbaiki data melalui tombol Perbarui Data Verifikasi, "
        "atau keluar dari aplikasi.",
        first_line_indent=Cm(1),
    )
    add_paragraph(
        doc,
        "Apabila admin meminta perbaikan atau menolak pengajuan, alumni dapat "
        "membaca catatan admin, memperbaiki data serta berkas melalui tombol "
        "Perbaiki Data & Kirim Ulang, kemudian mengirimkan kembali pengajuan "
        "verifikasi.",
        first_line_indent=Cm(1),
    )
    add_paragraph(
        doc,
        "Setelah status Terverifikasi, alumni secara otomatis diarahkan ke menu utama aplikasi.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "V. MENU UTAMA APLIKASI", 1)
    add_paragraph(
        doc,
        "Setelah akun terverifikasi, aplikasi menampilkan lima menu utama pada bagian bawah layar sebagai berikut.",
        first_line_indent=Cm(1),
    )
    add_table(
        doc,
        ["Menu", "Fungsi"],
        [
            ["Beranda", "Ringkasan layanan dan pintasan fitur"],
            ["Tracer", "Pengisian Tracer Study"],
            ["Loker", "Informasi dan lamaran lowongan kerja"],
            ["IKA", "Layanan Ikatan Alumni"],
            ["Profil", "Pengelolaan data diri"],
        ],
    )

    add_heading_custom(doc, "VI. BERANDA", 1)
    add_paragraph(
        doc,
        "Halaman Beranda menampilkan status layanan, menu pintasan, serta informasi penting "
        "berupa cuplikan pengumuman, lowongan, atau kegiatan terbaru.",
        first_line_indent=Cm(1),
    )
    add_paragraph(doc, "Pintasan yang tersedia pada Beranda:")
    add_table(
        doc,
        ["Menu", "Kegunaan"],
        [
            ["Informasi", "Pengumuman dan informasi umum"],
            ["Tracer", "Membuka Tracer Study"],
            ["Member IKA", "Membuka layanan IKA"],
            ["Loker", "Membuka lowongan kerja"],
            ["Event", "Daftar kegiatan alumni"],
            ["Digital Kartu Alumni", "Kartu alumni digital"],
            ["Lacak Teman", "Jejak Angkatan"],
            ["Notifikasi", "Pemberitahuan aplikasi"],
        ],
    )

    add_heading_custom(doc, "VII. TRACER STUDY", 1)
    add_paragraph(
        doc,
        "Tracer Study merupakan survei lulusan untuk mengetahui kondisi alumni setelah menyelesaikan studi.",
        first_line_indent=Cm(1),
    )
    add_paragraph(doc, "Langkah pengisian:")
    add_numbered(
        doc,
        [
            "Buka tab Tracer.",
            "Pilih Isi Tracer Study.",
            "Jawab seluruh pertanyaan dengan lengkap dan sesuai kondisi terkini.",
            "Simpan apabila belum selesai, atau kirim apabila sudah lengkap.",
            "Tinjau riwayat pengisian melalui menu Riwayat Pengisian.",
        ],
    )

    add_heading_custom(doc, "VIII. LOKER (LOWONGAN KERJA)", 1)
    add_paragraph(
        doc,
        "Fitur Loker menampilkan informasi lowongan kerja yang relevan bagi alumni.",
        first_line_indent=Cm(1),
    )
    add_paragraph(doc, "Langkah penggunaan:")
    add_numbered(
        doc,
        [
            "Buka tab Loker.",
            "Cari atau saring lowongan berdasarkan kategori, lokasi, atau tipe pekerjaan.",
            "Buka detail lowongan dan baca persyaratan.",
            "Lakukan lamaran sesuai petunjuk yang tertera, baik melalui aplikasi maupun saluran eksternal yang dicantumkan.",
        ],
    )
    add_paragraph(
        doc,
        "Kategori lowongan yang umum tersedia meliputi bidang keperawatan, kebidanan, "
        "tenaga kesehatan, klinik atau rumah sakit, puskesmas, dosen atau tutor, "
        "kerja luar negeri, studi lanjut, serta magang atau volunteer.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "IX. INFORMASI UMUM", 1)
    add_paragraph(
        doc,
        "Fitur Informasi Umum menampilkan pengumuman resmi terkait alumni, kampus, IKA, "
        "karier, kegiatan, dan informasi penting lainnya.",
        first_line_indent=Cm(1),
    )
    add_numbered(
        doc,
        [
            "Dari Beranda, pilih menu Informasi.",
            "Gunakan pencarian atau filter kategori sesuai kebutuhan.",
            "Buka detail pengumuman untuk membaca isi lengkap.",
            "Tandai pengumuman penting melalui fitur bookmark.",
        ],
    )

    add_heading_custom(doc, "X. EVENT ALUMNI", 1)
    add_paragraph(
        doc,
        "Fitur Event Alumni digunakan untuk mengikuti kegiatan alumni, seperti seminar, "
        "reuni, bakti sosial, dan pelatihan.",
        first_line_indent=Cm(1),
    )
    add_numbered(
        doc,
        [
            "Dari Beranda, pilih menu Event.",
            "Pilih kegiatan yang diminati.",
            "Baca detail waktu, lokasi, dan ketentuan kegiatan.",
            "Lakukan pendaftaran.",
            "Buka tiket kegiatan apabila telah disediakan.",
            "Periksa sertifikat atau galeri kegiatan setelah kegiatan berlangsung.",
        ],
    )
    add_paragraph(
        doc,
        "Kegiatan khusus anggota IKA tersedia pada menu IKA, sub-menu Event Anggota.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "XI. DIGITAL KARTU ALUMNI", 1)
    add_paragraph(
        doc,
        "Digital Kartu Alumni merupakan identitas digital bagi alumni yang telah terverifikasi. "
        "Kartu menampilkan data nama, NIM atau nomor alumni, program studi, angkatan, "
        "tahun lulus, dan kode QR.",
        first_line_indent=Cm(1),
    )
    add_numbered(
        doc,
        [
            "Dari Beranda, pilih menu Digital Kartu Alumni.",
            "Tinjau dan gunakan kartu sebagai identitas digital alumni.",
        ],
    )
    add_paragraph(
        doc,
        "Kartu Anggota IKA merupakan fitur terpisah dan hanya tersedia bagi anggota IKA aktif.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "XII. JEJAK ANGKATAN (LACAK TEMAN)", 1)
    add_paragraph(
        doc,
        "Fitur Jejak Angkatan memungkinkan alumni melihat daftar teman seangkatan atau "
        "seprogram studi yang mengizinkan profilnya ditampilkan.",
        first_line_indent=Cm(1),
    )
    add_numbered(
        doc,
        [
            "Dari Beranda, pilih menu Lacak Teman.",
            "Gunakan filter sesuai kebutuhan, misalnya satu kota, sudah bekerja, studi lanjut, atau anggota IKA.",
            "Buka profil teman untuk melihat informasi yang diizinkan tampil.",
        ],
    )
    add_paragraph(doc, "Ketentuan privasi:")
    add_bullets(
        doc,
        [
            "Nomor telepon tidak ditampilkan kepada alumni lain.",
            "Hanya data yang diizinkan pada pengaturan profil yang dapat dilihat.",
        ],
    )

    add_heading_custom(doc, "XIII. NOTIFIKASI", 1)
    add_paragraph(
        doc,
        "Fitur Notifikasi menyampaikan pemberitahuan penting terkait status verifikasi, "
        "Tracer Study, keanggotaan IKA, lowongan, kegiatan, pengumuman, serta pengingat iuran atau donasi.",
        first_line_indent=Cm(1),
    )
    add_numbered(
        doc,
        [
            "Dari Beranda, buka menu Notifikasi.",
            "Ketuk notifikasi untuk membaca detail.",
            "Tandai notifikasi sebagai telah dibaca.",
        ],
    )

    add_heading_custom(doc, "XIV. PROFIL ALUMNI", 1)
    add_paragraph(
        doc,
        "Pada tab Profil, alumni dapat mengelola data diri. Data yang dapat diperbarui meliputi "
        "foto profil, nomor telepon atau kontak, kota atau lokasi, status pekerjaan, instansi, "
        "jabatan, dan tautan media sosial.",
        first_line_indent=Cm(1),
    )
    add_paragraph(
        doc,
        "Data akademik seperti NIM, nama, program studi, angkatan, dan tahun lulus terkunci "
        "setelah proses verifikasi selesai. Apabila terdapat kekeliruan pada data akademik, "
        "alumni dapat menghubungi admin kampus untuk dilakukan perbaikan.",
        first_line_indent=Cm(1),
    )
    add_paragraph(doc, "Langkah menyimpan perubahan:")
    add_numbered(
        doc,
        [
            "Ubah data yang diperlukan.",
            "Ketuk tombol Simpan Perubahan.",
            "Pastikan data telah tersimpan sebelum meninggalkan halaman.",
        ],
    )

    add_heading_custom(doc, "XV. LAYANAN IKATAN ALUMNI (IKA)", 1)
    add_paragraph(
        doc,
        "IKA merupakan wadah Ikatan Alumni. Layanan IKA tersedia pada tab IKA.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "15.1 Status Keanggotaan", 2)
    add_paragraph(
        doc,
        "Pada beranda IKA, alumni dapat melihat status keanggotaan, antara lain: "
        "Belum Mendaftar, Menunggu Persetujuan, Ditolak, Anggota Aktif, atau Pengurus IKA.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "15.2 Pendaftaran IKA", 2)
    add_numbered(
        doc,
        [
            "Pastikan akun alumni telah berstatus Terverifikasi.",
            "Lengkapi nomor telepon pada tab Profil.",
            "Buka tab IKA.",
            "Pilih menu Daftar atau Ajukan Pendaftaran.",
            "Isi motivasi dan catatan sesuai formulir.",
            "Kirim pengajuan dan tunggu keputusan pengurus atau admin.",
        ],
    )
    add_paragraph(
        doc,
        "Setelah disetujui sebagai Anggota Aktif, fitur anggota IKA dapat digunakan.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "15.3 Fitur Anggota IKA Aktif", 2)
    add_table(
        doc,
        ["Fitur", "Kegunaan"],
        [
            ["Kartu Anggota", "Kartu digital anggota IKA beserta nomor anggota dan kode QR"],
            ["Event Anggota", "Kegiatan khusus anggota IKA, termasuk pendaftaran dan tiket"],
            ["Iuran / Donasi", "Melihat tagihan, mengunggah bukti pembayaran, dan riwayat"],
            ["Voting", "Pemberian suara pada pemilihan atau polling anggota"],
            ["Forum", "Ruang diskusi antar anggota"],
        ],
    )
    add_paragraph(
        doc,
        "Fitur di atas hanya dapat diakses oleh anggota IKA aktif.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "15.4 Iuran dan Donasi", 2)
    add_numbered(
        doc,
        [
            "Buka menu Iuran / Donasi.",
            "Pilih tagihan yang harus dibayar.",
            "Lakukan transfer sesuai instruksi pembayaran.",
            "Unggah bukti transfer.",
            "Tunggu konfirmasi dari pengurus atau admin.",
        ],
    )
    add_table(
        doc,
        ["Status", "Keterangan"],
        [
            ["Belum bayar", "Bukti pembayaran belum diunggah"],
            ["Menunggu konfirmasi", "Bukti pembayaran sedang diperiksa"],
            ["Lunas", "Pembayaran telah disetujui"],
            ["Ditolak", "Bukti ditolak; alumni mengunggah ulang sesuai catatan"],
        ],
    )

    add_heading_custom(doc, "15.5 Voting", 2)
    add_numbered(
        doc,
        [
            "Buka menu Voting.",
            "Pilih voting yang sedang berlangsung.",
            "Baca opsi yang tersedia.",
            "Pilih satu opsi dan kirim suara.",
            "Tinjau hasil voting apabila telah dipublikasikan.",
        ],
    )

    add_heading_custom(doc, "15.6 Forum", 2)
    add_numbered(
        doc,
        [
            "Buka menu Forum.",
            "Pilih kategori diskusi.",
            "Baca atau buat postingan.",
            "Berikan komentar atau tanggapan sesuai kebutuhan.",
        ],
    )

    add_heading_custom(doc, "XVI. RINGKASAN STATUS", 1)
    add_heading_custom(doc, "16.1 Status Verifikasi Alumni", 2)
    add_table(
        doc,
        ["Status", "Keterangan", "Tindak Lanjut"],
        [
            [
                "Menunggu Verifikasi",
                "Data sedang diperiksa admin",
                "Menunggu keputusan atau memperbarui data",
            ],
            [
                "Ditolak",
                "Pengajuan tidak disetujui",
                "Membaca alasan, memperbaiki data, dan mengirim ulang",
            ],
            [
                "Terverifikasi",
                "Akun disetujui",
                "Menggunakan menu utama aplikasi",
            ],
            [
                "Nonaktif",
                "Akun tidak aktif",
                "Menghubungi admin kampus",
            ],
        ],
    )

    add_heading_custom(doc, "16.2 Status Keanggotaan IKA", 2)
    add_table(
        doc,
        ["Status", "Keterangan"],
        [
            ["Belum Mendaftar", "Belum mengajukan keanggotaan"],
            ["Menunggu Persetujuan", "Pengajuan sedang diproses"],
            ["Ditolak", "Pengajuan tidak disetujui"],
            ["Anggota Aktif", "Dapat menggunakan fitur anggota"],
            ["Pengurus IKA", "Anggota sekaligus pengurus"],
        ],
    )

    add_heading_custom(doc, "XVII. PERTANYAAN YANG SERING DIAJUKAN", 1)

    faqs = [
        (
            "Mengapa saya belum dapat mengakses Beranda?",
            "Akun belum berstatus Terverifikasi. Periksa status pada layar verifikasi "
            "atau perbaiki data apabila diminta oleh admin.",
        ),
        (
            "Mengapa menu IKA tidak dapat dibuka?",
            "Status alumni terverifikasi berbeda dengan status anggota IKA. "
            "Alumni perlu mendaftar IKA dan menunggu persetujuan sebagai anggota aktif.",
        ),
        (
            "Kapan fitur anggota IKA dapat digunakan?",
            "Fitur anggota dapat digunakan setelah status keanggotaan menjadi Anggota Aktif.",
        ),
        (
            "Mengapa pendaftaran IKA gagal?",
            "Pastikan nomor telepon telah dilengkapi pada Profil, kemudian ajukan kembali.",
        ),
        (
            "Mengapa Informasi Umum kosong?",
            "Kondisi tersebut dapat terjadi apabila belum terdapat pengumuman yang dipublikasikan.",
        ),
        (
            "Apakah teman seangkatan dapat melihat foto ijazah saya?",
            "Tidak. Foto ijazah hanya digunakan untuk keperluan pemeriksaan oleh admin.",
        ),
        (
            "Apakah data akademik dapat diubah sendiri?",
            "Setelah terverifikasi, data akademik terkunci. Perbaikan dilakukan melalui admin kampus.",
        ),
        (
            "Mengapa suara pada Voting tidak dapat dikirim?",
            "Pastikan voting masih berlangsung, alumni belum pernah memberikan suara, "
            "dan salah satu opsi telah dipilih.",
        ),
        (
            "Apa yang harus dilakukan jika bukti iuran ditolak?",
            "Baca catatan penolakan, kemudian unggah ulang bukti pembayaran yang jelas dan sesuai ketentuan.",
        ),
        (
            "Bagaimana jika lupa kata sandi?",
            "Ikuti prosedur pemulihan akun sesuai ketentuan resmi kampus atau pengelola JejakGS.",
        ),
    ]
    for i, (q, a) in enumerate(faqs, start=1):
        add_paragraph(doc, f"{i}. {q}", bold=True, space_after=4)
        add_paragraph(doc, a, first_line_indent=Cm(0.5), space_after=10)

    add_heading_custom(doc, "XVIII. BANTUAN", 1)
    add_paragraph(
        doc,
        "Apabila mengalami kendala dalam penggunaan aplikasi, alumni disarankan untuk:",
        first_line_indent=Cm(1),
    )
    add_numbered(
        doc,
        [
            "Memastikan aplikasi telah diperbarui ke versi terbaru.",
            "Memastikan koneksi internet stabil.",
            "Mencoba memperbarui status atau masuk ulang ke aplikasi.",
            "Menghubungi pengelola layanan alumni STIKES Gunung Sari melalui kanal resmi kampus.",
        ],
    )
    add_paragraph(
        doc,
        "Saat menyampaikan laporan, mohon lampirkan nama lengkap, NIM, email akun JejakGS, "
        "ringkasan permasalahan, serta tangkapan layar apabila tersedia.",
        first_line_indent=Cm(1),
    )

    add_heading_custom(doc, "XIX. RINGKASAN ALUR PENGGUNAAN", 1)
    add_paragraph(
        doc,
        "Secara ringkas, alur penggunaan aplikasi JejakGS adalah sebagai berikut:",
        first_line_indent=Cm(1),
    )
    add_numbered(
        doc,
        [
            "Mengunduh dan membuka aplikasi JejakGS.",
            "Melakukan pendaftaran atau masuk ke akun.",
            "Melengkapi profil dan mengunggah foto ijazah serta foto diri.",
            "Menunggu proses verifikasi oleh admin.",
            "Memperbaiki dan mengirim ulang data apabila diminta.",
            "Menggunakan fitur layanan alumni setelah status Terverifikasi.",
            "Mendaftar IKA secara opsional untuk mengakses layanan anggota.",
            "Menggunakan fitur IKA setelah berstatus Anggota Aktif.",
        ],
    )

    add_horizontal_line(doc)
    add_paragraph(
        doc,
        "Dokumen ini disusun sebagai panduan penggunaan aplikasi JejakGS bagi alumni "
        "STIKES Gunung Sari. Segala ketentuan terkait kebijakan verifikasi, iuran, "
        "dan kegiatan IKA mengikuti ketentuan resmi STIKES Gunung Sari serta pengurus IKA.",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
